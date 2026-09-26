"""Predict ETA and delivery risk for active Zamato orders.

Learns delivery time from historical orders (Capstone cleaned data), then scores
the live feed in active_deliveries_ZC.xlsx:

- Predicted_ETA_min  : LightGBM regression trained on historical Time_taken (min)
- Risk_Score         : the team's historical risk rule (weather + traffic +
                       festival + multiple deliveries, capped at 100)
- Risk_Category      : Low / Medium / High / Critical
- Weather / traffic / festival labels, and a High+Critical vs. rest split

Usage:
    python src/predict_eta.py
"""
from datetime import datetime, time
from pathlib import Path

import joblib
import lightgbm as lgb
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import train_test_split

ROOT = Path(__file__).resolve().parents[1]
HIST_PATH = ROOT / "data/raw/Capstone_Team1_CleanedData_v1_Zamato_Delivery_Operation.xlsx"
HIST_SHEET = "Zamato Delivery Filtering"
ACTIVE_PATH = ROOT / "data/raw/active_deliveries_ZC.xlsx"
OUT_DIR = ROOT / "outputs"
MODEL_PATH = ROOT / "models/eta_lgbm.joblib"
RANDOM_STATE = 42

# ---------------------------------------------------------------------------
# Risk rules. Reproduces the historical "Risk Score" column exactly
# (validated on all 40,088 historical rows).
# ---------------------------------------------------------------------------
WEATHER_POINTS = {"Sunny": 0, "Windy": 10, "Cloudy": 20, "Fog": 20, "Sandstorms": 20, "Stormy": 20}
TRAFFIC_POINTS = {"Low": 0, "Medium": 10, "High": 20, "Jam": 30}
FESTIVAL_POINTS = {"No": 0, "Yes": 30}
MULTI_DELIVERY_POINTS_PER_ORDER = 10
RISK_CAP = 100
# Category thresholds (lower bound, inclusive). Chosen from historical data:
# score >= 70 -> 57-95% of orders took > 40 min; 50-60 -> avg 30-33 min.
RISK_THRESHOLDS = [("Critical", 70), ("High", 50), ("Medium", 30), ("Low", 0)]

WEATHER_LABELS = {"Sunny": "Clear", "Windy": "Moderate", "Cloudy": "Adverse", "Fog": "Adverse",
                  "Sandstorms": "Adverse", "Stormy": "Adverse"}
TRAFFIC_LABELS = {"Low": "Smooth", "Medium": "Moderate", "High": "Heavy", "Jam": "Severe"}
FESTIVAL_LABELS = {"No": "Normal Day", "Yes": "Festival / Event"}

CATEGORICAL = ["Weather_conditions", "Road_traffic_density", "Type_of_order",
               "Type_of_vehicle", "Festival", "City"]
NUMERIC = ["Delivery_person_Age", "Delivery_person_Ratings", "Vehicle_condition",
           "multiple_deliveries", "distance_km", "pickup_wait_min", "order_hour"]
FEATURES = NUMERIC + CATEGORICAL
MATCH_KEYS = ["Delivery_person_ID", "Restaurant_latitude", "Restaurant_longitude",
              "Delivery_location_latitude", "Delivery_location_longitude", "Weather_conditions",
              "Road_traffic_density", "Type_of_order"]


def _to_minutes(value):
    """Minutes since midnight from a time / datetime / 'HH:MM:SS' string."""
    if pd.isna(value):
        return np.nan
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, str):
        value = time.fromisoformat(value.strip())
    return value.hour * 60 + value.minute


def haversine_km(lat1, lon1, lat2, lon2):
    lat1, lon1, lat2, lon2 = map(np.radians, (lat1, lon1, lat2, lon2))
    a = np.sin((lat2 - lat1) / 2) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2
    return 6371 * 2 * np.arcsin(np.sqrt(a))


def build_features(df):
    out = df.copy()
    # Some restaurant coordinates are sign-flipped in the source; delivery ones are not.
    out["distance_km"] = haversine_km(out["Restaurant_latitude"].abs(), out["Restaurant_longitude"].abs(),
                                      out["Delivery_location_latitude"].abs(),
                                      out["Delivery_location_longitude"].abs())
    ordered = out["Time_Ordered"].map(_to_minutes)
    picked = out["Time_Order_picked"].map(_to_minutes)
    out["pickup_wait_min"] = (picked - ordered) % (24 * 60)
    out["order_hour"] = ordered // 60
    for col in CATEGORICAL:
        out[col] = pd.Categorical(out[col].astype(str), categories=CATEGORY_LEVELS[col])
    return out


def add_risk(df):
    out = df.copy()
    out["Weather_Label"] = out["Weather_conditions"].map(WEATHER_LABELS)
    out["Traffic_Label"] = out["Road_traffic_density"].map(TRAFFIC_LABELS)
    out["Festival_Label"] = out["Festival"].map(FESTIVAL_LABELS)
    out["Weather_Pts"] = out["Weather_conditions"].map(WEATHER_POINTS)
    out["Traffic_Pts"] = out["Road_traffic_density"].map(TRAFFIC_POINTS)
    out["Festival_Pts"] = out["Festival"].map(FESTIVAL_POINTS)
    out["MultiDelivery_Pts"] = out["multiple_deliveries"] * MULTI_DELIVERY_POINTS_PER_ORDER
    pts = ["Weather_Pts", "Traffic_Pts", "Festival_Pts", "MultiDelivery_Pts"]
    out["Risk_Score"] = out[pts].sum(axis=1).clip(upper=RISK_CAP)
    out["Risk_Category"] = out["Risk_Score"].map(
        lambda s: next(name for name, lo in RISK_THRESHOLDS if s >= lo))

    def drivers(r):
        items = []
        if r.Weather_Pts:
            items.append(f"Weather: {r.Weather_conditions}")
        if r.Traffic_Pts:
            items.append(f"Traffic: {r.Road_traffic_density}")
        if r.Festival_Pts:
            items.append("Festival")
        if r.MultiDelivery_Pts:
            items.append(f"Multiple deliveries: {int(r.multiple_deliveries)}")
        return "; ".join(items) or "None"

    out["Risk_Drivers"] = out.apply(drivers, axis=1)
    return out


# ---------------------------------------------------------------------------
# Load data
# ---------------------------------------------------------------------------
hist = pd.read_excel(HIST_PATH, sheet_name=HIST_SHEET).dropna(subset=["ID"])
active_book = pd.read_excel(ACTIVE_PATH, sheet_name=None)
active = active_book["Active_Deliveries"]
truth = active_book["Ground_Truth"]

CATEGORY_LEVELS = {c: sorted(hist[c].dropna().astype(str).unique()) for c in CATEGORICAL}

# Sanity check: our rule reproduces the historical Risk Score.
rule_match = (add_risk(hist)["Risk_Score"] == hist["Risk Score"]).mean()
assert rule_match == 1.0, f"Risk rule reproduces only {rule_match:.2%} of history"

# The active batch was sampled from history; drop those rows so the backtest is honest.
merge_hist = hist[MATCH_KEYS].round(6).astype(str)
merge_act = active[MATCH_KEYS].round(6).astype(str)
overlap = merge_hist.merge(merge_act.drop_duplicates(), on=MATCH_KEYS, how="left", indicator=True)
hist_train = hist[(overlap["_merge"] == "left_only").to_numpy()]
n_removed = len(hist) - len(hist_train)

# ---------------------------------------------------------------------------
# Train & evaluate
# ---------------------------------------------------------------------------
Xh = build_features(hist_train)[FEATURES]
yh = hist_train["Time_taken (min)"]
X_tr, X_te, y_tr, y_te = train_test_split(Xh, yh, test_size=0.2, random_state=RANDOM_STATE)

params = dict(n_estimators=600, learning_rate=0.05, num_leaves=63, min_child_samples=30,
              subsample=0.8, subsample_freq=1, colsample_bytree=0.8,
              random_state=RANDOM_STATE, verbose=-1)
model = lgb.LGBMRegressor(**params).fit(X_tr, y_tr)
pred_te = model.predict(X_te)
holdout = {
    "MAE (min)": mean_absolute_error(y_te, pred_te),
    "RMSE (min)": mean_squared_error(y_te, pred_te) ** 0.5,
    "R2": r2_score(y_te, pred_te),
}
baseline_mae = mean_absolute_error(y_te, np.full(len(y_te), y_tr.mean()))

# Refit on all historical data for the live predictions.
model = lgb.LGBMRegressor(**params).fit(Xh, yh)
MODEL_PATH.parent.mkdir(exist_ok=True)
joblib.dump({"model": model, "features": FEATURES, "category_levels": CATEGORY_LEVELS}, MODEL_PATH)
importance = (pd.Series(model.booster_.feature_importance("gain"), index=FEATURES)
              .sort_values(ascending=False))
importance = importance / importance.sum()

# ---------------------------------------------------------------------------
# Score the active feed
# ---------------------------------------------------------------------------
run_ts = pd.Timestamp.now().floor("s").to_pydatetime()
Xa = build_features(active)
scored = add_risk(active)
scored["Predicted_ETA_min"] = np.round(model.predict(Xa[FEATURES]), 1)
scored["ETA_Updated_At"] = run_ts
scored["Distance_km"] = Xa["distance_km"].round(2)
ordered_dt = pd.to_datetime(scored["Order_Date"].astype(str) + " " + scored["Time_Ordered"].astype(str))
scored["Estimated_Arrival"] = (ordered_dt + pd.to_timedelta(scored["Predicted_ETA_min"], unit="m")).dt.strftime("%H:%M")

backtest = scored[["Order_ID", "Predicted_ETA_min", "Risk_Category"]].merge(truth, on="Order_ID")
bt_mae = mean_absolute_error(backtest["Actual_Time_taken_min"], backtest["Predicted_ETA_min"])

# ---------------------------------------------------------------------------
# Write workbook
# ---------------------------------------------------------------------------
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
NEW_FILL = PatternFill("solid", fgColor="DDEBF7")
BODY_FONT = Font(name=FONT)
INPUT_FONT = Font(name=FONT, color="0000FF")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
CAT_FILLS = {"Critical": "F8CBAD", "High": "FFE699", "Medium": "E2EFDA", "Low": "FFFFFF"}

ORIGINAL_COLS = list(active.columns)
ADDED_COLS = ["Weather_Label", "Traffic_Label", "Festival_Label", "Weather_Pts", "Traffic_Pts",
              "Festival_Pts", "MultiDelivery_Pts", "Risk_Drivers", "Distance_km", "Estimated_Arrival"]
SHEET_COLS = ORIGINAL_COLS + ADDED_COLS


def style_header(ws, row, cols, highlight=()):
    for j, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if name in highlight:
            c.fill = PatternFill("solid", fgColor="2E75B6")
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def autosize(ws, min_w=9, max_w=34):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[letter].width = max(min_w, min(max_w, width + 2))


def write_value(cell, value):
    if isinstance(value, (np.integer,)):
        value = int(value)
    elif isinstance(value, (np.floating,)):
        value = None if np.isnan(value) else float(value)
    cell.value = value
    cell.font = BODY_FONT


def write_table(ws, df, cols, start_row=1, highlight=()):
    style_header(ws, start_row, cols, highlight)
    for i, rec in enumerate(df[cols].itertuples(index=False), start_row + 1):
        cat = rec[cols.index("Risk_Category")] if "Risk_Category" in cols else None
        for j, v in enumerate(rec, 1):
            c = ws.cell(row=i, column=j)
            write_value(c, v)
            c.border = THIN
            if cat and cols[j - 1] in ("Risk_Score", "Risk_Category"):
                c.fill = PatternFill("solid", fgColor=CAT_FILLS[cat])
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(df)}"


wb = Workbook()

# --- Risk_Rules (inputs referenced by formulas) ----------------------------
rules = wb.active
rules.title = "Risk_Rules"
rules["A1"] = "Risk scoring rules (edit blue cells; Active_Deliveries recalculates)"
rules["A1"].font = Font(name=FONT, bold=True, size=13)
rules["A2"] = ("Source: reverse-engineered from the historical 'Risk Score' column in "
               "Capstone_Team1_CleanedData_v1 (matches 40,088 / 40,088 rows).")
rules["A2"].font = Font(name=FONT, italic=True, color="595959")


def rule_block(top, title, mapping, label_map=None):
    rules.cell(row=top, column=1, value=title).font = Font(name=FONT, bold=True)
    hdr = ["Value", "Points"] + (["Label"] if label_map else [])
    for j, h in enumerate(hdr, 1):
        c = rules.cell(row=top + 1, column=j, value=h)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
    for i, (k, pts) in enumerate(mapping.items(), top + 2):
        rules.cell(row=i, column=1, value=k).font = BODY_FONT
        rules.cell(row=i, column=2, value=pts).font = INPUT_FONT
        if label_map:
            rules.cell(row=i, column=3, value=label_map[k]).font = BODY_FONT
    first, last = top + 2, top + 1 + len(mapping)
    return f"Risk_Rules!$A${first}:$A${last}", f"Risk_Rules!$B${first}:$B${last}", top + 3 + len(mapping)


r = 4
W_KEYS, W_PTS, r = rule_block(r, "Weather conditions", WEATHER_POINTS, WEATHER_LABELS)
T_KEYS, T_PTS, r = rule_block(r, "Road traffic density", TRAFFIC_POINTS, TRAFFIC_LABELS)
F_KEYS, F_PTS, r = rule_block(r, "Festival / event", FESTIVAL_POINTS, FESTIVAL_LABELS)
rules.cell(row=r, column=1, value="Other parameters").font = Font(name=FONT, bold=True)
rules.cell(row=r + 1, column=1, value="Points per extra delivery (multiple_deliveries)").font = BODY_FONT
rules.cell(row=r + 1, column=2, value=MULTI_DELIVERY_POINTS_PER_ORDER).font = INPUT_FONT
rules.cell(row=r + 2, column=1, value="Risk score cap").font = BODY_FONT
rules.cell(row=r + 2, column=2, value=RISK_CAP).font = INPUT_FONT
MD_REF, CAP_REF = f"Risk_Rules!$B${r + 1}", f"Risk_Rules!$B${r + 2}"
r += 4
rules.cell(row=r, column=1, value="Risk category thresholds (score >= lower bound)").font = Font(name=FONT, bold=True)
for j, h in enumerate(["Category", "Lower bound", "Historical evidence"], 1):
    c = rules.cell(row=r + 1, column=j, value=h)
    c.font, c.fill = HEADER_FONT, HEADER_FILL
evidence = {"Critical": "Score 70-100: avg 42-47 min, 57-95% of orders > 40 min",
            "High": "Score 50-60: avg 30-33 min, 11-19% of orders > 40 min",
            "Medium": "Score 30-40: avg 23-27 min, 3-6% of orders > 40 min",
            "Low": "Score 0-20: avg 19-21 min, ~1% of orders > 40 min"}
THRESH = {}
for i, (name, lo) in enumerate(RISK_THRESHOLDS, r + 2):
    rules.cell(row=i, column=1, value=name).font = BODY_FONT
    rules.cell(row=i, column=2, value=lo).font = INPUT_FONT
    rules.cell(row=i, column=3, value=evidence[name]).font = BODY_FONT
    THRESH[name] = f"Risk_Rules!$B${i}"
rules.column_dimensions["A"].width = 46
rules.column_dimensions["B"].width = 14
rules.column_dimensions["C"].width = 58

# --- Active_Deliveries (full feed, formulas for risk) -----------------------
ws = wb.create_sheet("Active_Deliveries", 0)
full = scored.reset_index(drop=True)  # keep live-feed order
write_table(ws, full, SHEET_COLS,
            highlight={"Predicted_ETA_min", "ETA_Updated_At", "Risk_Score", "Risk_Category"} | set(ADDED_COLS))
col = {name: get_column_letter(i) for i, name in enumerate(SHEET_COLS, 1)}
for i in range(2, len(full) + 2):
    ws[f"{col['Weather_Pts']}{i}"] = f"=INDEX({W_PTS},MATCH({col['Weather_conditions']}{i},{W_KEYS},0))"
    ws[f"{col['Traffic_Pts']}{i}"] = f"=INDEX({T_PTS},MATCH({col['Road_traffic_density']}{i},{T_KEYS},0))"
    ws[f"{col['Festival_Pts']}{i}"] = f"=INDEX({F_PTS},MATCH({col['Festival']}{i},{F_KEYS},0))"
    ws[f"{col['MultiDelivery_Pts']}{i}"] = f"={col['multiple_deliveries']}{i}*{MD_REF}"
    ws[f"{col['Risk_Score']}{i}"] = (f"=MIN({CAP_REF},{col['Weather_Pts']}{i}+{col['Traffic_Pts']}{i}"
                                     f"+{col['Festival_Pts']}{i}+{col['MultiDelivery_Pts']}{i})")
    s = f"{col['Risk_Score']}{i}"
    ws[f"{col['Risk_Category']}{i}"] = (f'=IF({s}>={THRESH["Critical"]},"Critical",IF({s}>={THRESH["High"]},'
                                        f'"High",IF({s}>={THRESH["Medium"]},"Medium","Low")))')
    for name in ("Weather_Pts", "Traffic_Pts", "Festival_Pts", "MultiDelivery_Pts", "Risk_Score", "Risk_Category"):
        ws[f"{col[name]}{i}"].font = BODY_FONT
    ws[f"{col['ETA_Updated_At']}{i}"].number_format = "yyyy-mm-dd hh:mm:ss"
ws[f"{col['Predicted_ETA_min']}1"].comment = Comment(
    "LightGBM model trained on historical Zamato deliveries (minutes from order to delivery).", "ETA model")
ws[f"{col['Risk_Score']}1"].comment = Comment(
    "Formula: MIN(cap, weather + traffic + festival + multi-delivery points). Rules in Risk_Rules.", "ETA model")
ws[f"{col['Ops_Alerted']}1"].comment = Comment(
    "Left blank: this project only scores orders; no alerts or notifications were sent.", "ETA model")
autosize(ws)

# --- Split sheets (static values, same rules as above) ----------------------
is_hc = scored["Risk_Category"].isin(["High", "Critical"])
order = ["Risk_Score", "Predicted_ETA_min"]
split_cols = ["Order_ID", "Delivery_person_ID", "City", "Time_Ordered", "Time_Order_picked",
              "Weather_conditions", "Weather_Label", "Road_traffic_density", "Traffic_Label",
              "Festival", "Festival_Label", "multiple_deliveries", "Type_of_vehicle", "Distance_km",
              "Predicted_ETA_min", "Estimated_Arrival", "Risk_Score", "Risk_Category", "Risk_Drivers"]
for title, part in [("High_Critical", scored[is_hc]), ("Low_Medium", scored[~is_hc])]:
    sh = wb.create_sheet(title)
    sh["A1"] = (f"{title.replace('_', ' + ')} risk orders ({len(part)} of {len(scored)}), "
                "sorted by risk score then ETA. Values snapshot from the model run.")
    sh["A1"].font = Font(name=FONT, bold=True)
    write_table(sh, part.sort_values(order, ascending=False), split_cols, start_row=3)
    sh.freeze_panes = "B4"
    autosize(sh)

# --- Backtest -----------------------------------------------------------------
bt = wb.create_sheet("Backtest")
bt["A1"] = "Backtest vs. Ground_Truth (actual time taken). Not used for training."
bt["A1"].font = Font(name=FONT, bold=True)
bt_cols = ["Order_ID", "Risk_Category", "Predicted_ETA_min", "Actual_Time_taken_min", "Error_min", "Abs_Error_min"]
style_header(bt, 3, bt_cols)
for i, rec in enumerate(backtest.itertuples(index=False), 4):
    for j, v in enumerate([rec.Order_ID, rec.Risk_Category, rec.Predicted_ETA_min, rec.Actual_Time_taken_min], 1):
        write_value(bt.cell(row=i, column=j), v)
    bt[f"E{i}"] = f"=C{i}-D{i}"
    bt[f"F{i}"] = f"=ABS(E{i})"
    bt[f"E{i}"].font = bt[f"F{i}"].font = BODY_FONT
last = 3 + len(backtest)
for k, (label, formula) in enumerate([
        ("MAE (min)", f"=AVERAGE(F4:F{last})"),
        ("RMSE (min)", f"=SQRT(SUMPRODUCT(E4:E{last},E4:E{last})/COUNT(E4:E{last}))"),
        ("Mean error / bias (min)", f"=AVERAGE(E4:E{last})"),
        ("Orders within 5 min", f'=COUNTIF(F4:F{last},"<=5")/COUNT(F4:F{last})')], 0):
    bt.cell(row=3 + k, column=8, value=label).font = Font(name=FONT, bold=True)
    c = bt.cell(row=3 + k, column=9, value=formula)
    c.font = BODY_FONT
    c.number_format = "0.0%" if "within" in label else "0.00"
autosize(bt)
bt.column_dimensions["H"].width = 26

# --- Summary ------------------------------------------------------------------
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Active Deliveries (batch ZC): ETA and Risk Summary"
sm["A1"].font = Font(name=FONT, bold=True, size=14)
sm["A2"] = f"Model run: {run_ts:%Y-%m-%d %H:%M:%S}  |  Orders scored: {len(scored)}"
sm["A2"].font = Font(name=FONT, italic=True, color="595959")
for j, h in enumerate(["Risk category", "Orders", "Share", "Avg predicted ETA (min)", "Max predicted ETA (min)"], 1):
    c = sm.cell(row=4, column=j, value=h)
    c.font, c.fill = HEADER_FONT, HEADER_FILL
    c.alignment = Alignment(wrap_text=True, horizontal="center")
AD = "Active_Deliveries"
cat_rng = f"{AD}!${col['Risk_Category']}$2:${col['Risk_Category']}${len(full) + 1}"
eta_rng = f"{AD}!${col['Predicted_ETA_min']}$2:${col['Predicted_ETA_min']}${len(full) + 1}"
for i, (name, _) in enumerate(RISK_THRESHOLDS, 5):
    sm.cell(row=i, column=1, value=name).font = Font(name=FONT, bold=True)
    sm.cell(row=i, column=1).fill = PatternFill("solid", fgColor=CAT_FILLS[name])
    sm[f"B{i}"] = f'=COUNTIF({cat_rng},A{i})'
    sm[f"C{i}"] = f"=IF($B$9=0,0,B{i}/$B$9)"
    sm[f"D{i}"] = f'=IF(B{i}=0,"-",AVERAGEIF({cat_rng},A{i},{eta_rng}))'
    sm[f"E{i}"] = f'=IF(B{i}=0,"-",_xlfn.MAXIFS({eta_rng},{cat_rng},A{i}))'
sm["A9"] = "Total"
sm["B9"] = "=SUM(B5:B8)"
sm["C9"] = "=SUM(C5:C8)"
sm["D9"] = f"=AVERAGE({eta_rng})"
sm["E9"] = f"=MAX({eta_rng})"
sm["A10"] = "High + Critical (needs attention)"
sm["B10"] = "=B5+B6"
sm["C10"] = "=IF($B$9=0,0,B10/$B$9)"
for rr in range(5, 11):
    for cc in "ABCDE":
        sm[f"{cc}{rr}"].font = Font(name=FONT, bold=rr >= 9)
    sm[f"C{rr}"].number_format = "0.0%"
    sm[f"D{rr}"].number_format = sm[f"E{rr}"].number_format = "0.0"

sm["A12"] = "Label breakdown"
sm["A12"].font = Font(name=FONT, bold=True, size=12)
row = 13
for field, labels in [("Weather_Label", ["Clear", "Moderate", "Adverse"]),
                      ("Traffic_Label", ["Smooth", "Moderate", "Heavy", "Severe"]),
                      ("Festival_Label", ["Normal Day", "Festival / Event"])]:
    rng = f"{AD}!${col[field]}$2:${col[field]}${len(full) + 1}"
    for lab in labels:
        sm.cell(row=row, column=1, value=f"{field.replace('_Label', '')}: {lab}").font = BODY_FONT
        sm[f"B{row}"] = f'=COUNTIF({rng},"{lab}")'
        sm[f"B{row}"].font = BODY_FONT
        row += 1

row += 1
sm.cell(row=row, column=1, value="ETA model (LightGBM)").font = Font(name=FONT, bold=True, size=12)
model_rows = [
    ("Training rows (historical, valid)", len(hist_train)),
    ("Historical rows removed (same orders as active batch)", n_removed),
    ("Hold-out MAE (min)", round(holdout["MAE (min)"], 2)),
    ("Hold-out RMSE (min)", round(holdout["RMSE (min)"], 2)),
    ("Hold-out R2", round(holdout["R2"], 3)),
    ("Baseline MAE: predict the mean (min)", round(baseline_mae, 2)),
    ("Backtest MAE on active batch (min)", "=Backtest!I3"),
]
for k, (label, val) in enumerate(model_rows, row + 1):
    sm.cell(row=k, column=1, value=label).font = BODY_FONT
    c = sm.cell(row=k, column=2, value=val)
    c.font = BODY_FONT
    if label.startswith("Backtest"):
        c.number_format = "0.00"
row = row + len(model_rows) + 2
sm.cell(row=row, column=1, value="Top features (share of total gain)").font = Font(name=FONT, bold=True)
for k, (feat, share) in enumerate(importance.head(8).items(), row + 1):
    sm.cell(row=k, column=1, value=feat).font = BODY_FONT
    c = sm.cell(row=k, column=2, value=float(share))
    c.font, c.number_format = BODY_FONT, "0.0%"
sm.column_dimensions["A"].width = 52
for cc in "BCDE":
    sm.column_dimensions[cc].width = 16

OUT_DIR.mkdir(exist_ok=True)
out_xlsx = OUT_DIR / "active_deliveries_ZC_predicted.xlsx"
wb.save(out_xlsx)

# CSV copies for the dashboard.
csv_cols = SHEET_COLS
scored[csv_cols].to_csv(OUT_DIR / "active_deliveries_ZC_predicted.csv", index=False)
scored[is_hc][split_cols].sort_values(order, ascending=False).to_csv(OUT_DIR / "high_critical_orders.csv", index=False)
scored[~is_hc][split_cols].sort_values(order, ascending=False).to_csv(OUT_DIR / "low_medium_orders.csv", index=False)

print(f"Removed {n_removed} overlapping historical rows; trained on {len(hist_train)}")
print("Hold-out:", {k: round(v, 3) for k, v in holdout.items()}, "| baseline MAE", round(baseline_mae, 2))
print(f"Backtest MAE on active batch: {bt_mae:.2f} min")
print(scored["Risk_Category"].value_counts().to_dict())
print(f"Saved {out_xlsx.relative_to(ROOT)}")
